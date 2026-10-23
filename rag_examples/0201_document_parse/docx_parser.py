from pathlib import Path

from docx import Document


PROJECT_DIR = Path(__file__).resolve().parents[2]
DATA_DIR = PROJECT_DIR / "rag_examples" / "data"
OUTPUT_DIR = PROJECT_DIR / "rag_examples" / "data" / "output"


def save_images(document, file_path, output_dir):
    """提取 DOCX 内嵌图片并返回保存数量。"""
    image_dir = output_dir / file_path.stem / "images"
    count = 0

    for part in document.part.package.parts:
        if not str(part.partname).startswith("/word/media/"):
            continue
        image_dir.mkdir(parents=True, exist_ok=True)
        (image_dir / Path(str(part.partname)).name).write_bytes(part.blob)
        count += 1
    return count


def parse_docx(file_path, output_dir=OUTPUT_DIR):
    """返回 DOCX 文本，并将内嵌图片保存到 output。"""
    file_path = Path(file_path)
    document = Document(str(file_path))
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    # 将表格的每一行转换为制表符分隔文本
    for table in document.tables:
        for row in table.rows:
            text = "\t".join(cell.text.strip().replace("\n", " ") for cell in row.cells)
            if text.strip():
                lines.append(text)

    save_images(document, file_path, output_dir)
    return "\n".join(lines)


def main():
    # 批量测试 datas 目录下的 DOCX 文件
    for file_path in DATA_DIR.rglob("*.docx"):
        text = parse_docx(file_path)
        print(text[:100])
        print("DOCX 解析完成：{}，共 {} 个字符".format(file_path.name, len(text)))


if __name__ == "__main__":
    main()
